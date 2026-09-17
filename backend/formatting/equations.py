"""
Equation Formatting Engine.

Converts mathematical expressions and LaTeX code into native Word OMML
(Office Math Markup Language) elements using latex2mathml and Microsoft's
MML2OMML stylesheet.
"""

from __future__ import annotations
import os
import re
from typing import Optional, List
import docx
from docx.oxml import parse_xml
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import latex2mathml.converter
from lxml import etree

_XSLT_PATH = os.path.join(os.path.dirname(__file__), "..", "export", "MML2OMML.XSL")
_TRANSFORM = None


def _get_transform():
    global _TRANSFORM
    if _TRANSFORM is None:
        if os.path.exists(_XSLT_PATH):
            xslt_doc = etree.parse(_XSLT_PATH)
            _TRANSFORM = etree.XSLT(xslt_doc)
        else:
            raise FileNotFoundError(f"MML2OMML.XSL not found at {_XSLT_PATH}")
    return _TRANSFORM


def latex_to_omml(latex_code: str, display: bool = False):
    """
    Converts a LaTeX equation string to an OpenXML OMML element.
    Returns a docx OxmlElement (<m:oMath> or <m:oMathPara>) or None if conversion fails.
    """
    cleaned = latex_code.strip()
    if cleaned.startswith("$$") and cleaned.endswith("$$"):
        cleaned = cleaned[2:-2].strip()
        display = True
    elif cleaned.startswith("$") and cleaned.endswith("$"):
        cleaned = cleaned[1:-1].strip()

    if not cleaned:
        return None

    # Normalizations for common LaTeX edge cases
    cleaned = cleaned.replace(r"\left(", "(").replace(r"\right)", ")")
    cleaned = cleaned.replace(r"\left[", "[").replace(r"\right]", "]")
    cleaned = cleaned.replace(r"\left.", "").replace(r"\right.", "")
    cleaned = re.sub(r",\s*de\b", r"\,de", cleaned)

    try:
        mathml = latex2mathml.converter.convert(cleaned)
        transform = _get_transform()
        omml_tree = transform(etree.fromstring(mathml.encode("utf-8")))
        omml_str = etree.tostring(omml_tree, encoding="utf-8").decode("utf-8")

        if display:
            # Wrap in oMathPara for display equations if not already wrapped
            if not omml_str.startswith("<m:oMathPara") and not omml_str.startswith("<oMathPara"):
                omml_str = (
                    f'<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                    f"{omml_str}</m:oMathPara>"
                )
        return parse_xml(omml_str)
    except Exception:
        return None


def is_equation_text(text: str) -> bool:
    """Checks if text contains LaTeX equation markers."""
    t = text.strip()
    if t.startswith("$$") and t.endswith("$$"):
        return True
    if "$" in t:
        return True
    return False


def render_paragraph_content(
    p: docx.text.paragraph.Paragraph,
    text: str,
    runs: Optional[List] = None,
    font_name: str = "Times New Roman",
    font_size_pt: float = 11.0,
    is_bold: bool = False,
    is_italic: bool = False,
    is_display_equation: bool = False,
) -> None:
    """
    Renders paragraph content, seamlessly converting any display or inline
    LaTeX equations into native Word OMML.
    """
    clean_text = text.strip()

    # Case 1: Standalone Display Equation
    if is_display_equation or (clean_text.startswith("$$") and clean_text.endswith("$$")):
        eq_element = latex_to_omml(clean_text, display=True)
        if eq_element is not None:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p._p.append(eq_element)
            return

    # Case 2: Contains Inline Math ($...$)
    if "$" in text:
        parts = re.split(r"(\$[^\$]+?\$)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("$") and part.endswith("$") and len(part) > 2:
                # Math expression
                math_code = part[1:-1]
                omml_el = latex_to_omml(math_code, display=False)
                if omml_el is not None:
                    p._p.append(omml_el)
                    continue
                else:
                    # Fallback to italic run if OMML fails
                    r = p.add_run(math_code)
                    r.font.name = font_name
                    r.font.size = Pt(font_size_pt)
                    r.italic = True
                    continue
            # Normal text run
            r = p.add_run(part)
            r.font.name = font_name
            r.font.size = Pt(font_size_pt)
            r.bold = is_bold
            r.italic = is_italic
        return

    # Case 3: Standard runs preservation if provided
    if runs:
        for r_meta in runs:
            r = p.add_run(r_meta.text)
            r.font.name = r_meta.font_name or font_name
            if r_meta.font_size_pt:
                r.font.size = Pt(r_meta.font_size_pt)
            else:
                r.font.size = Pt(font_size_pt)
            r.bold = r_meta.bold if r_meta.bold is not None else is_bold
            r.italic = r_meta.italic if r_meta.italic is not None else is_italic
            r.underline = r_meta.underline
            if r_meta.color and r_meta.color.startswith("#") and len(r_meta.color) == 7:
                try:
                    hex_val = r_meta.color.lstrip("#")
                    r.font.color.rgb = RGBColor(
                        int(hex_val[0:2], 16),
                        int(hex_val[2:4], 16),
                        int(hex_val[4:6], 16),
                    )
                except Exception:
                    pass
        return

    # Case 4: Default single run
    r = p.add_run(text)
    r.font.name = font_name
    r.font.size = Pt(font_size_pt)
    r.bold = is_bold
    r.italic = is_italic
