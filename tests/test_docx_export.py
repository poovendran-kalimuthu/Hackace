"""
Unit Tests for Document Quality Validation and DOCX Export.
"""

import os
import pytest
import docx

from backend.document.model import Block, BlockType, DocumentModel
from backend.export.docx import DocxExporter
from backend.export.validator import DocumentQualityValidator
from backend.templates.manager import TemplateManager


@pytest.fixture
def classic_template():
    mgr = TemplateManager()
    return mgr.get_template("classic_novel")


def test_quality_validator_passes():
    doc = DocumentModel(
        blocks=[
            Block(block_type=BlockType.CHAPTER_TITLE, text="Chapter 1"),
            Block(block_type=BlockType.BODY, text="A valid body paragraph."),
        ]
    )
    result = DocumentQualityValidator.validate(doc)
    assert result.is_valid
    assert len(result.errors) == 0


def test_docx_exporter_creates_valid_file(tmp_path, classic_template):
    out_file = str(tmp_path / "published_book.docx")
    exporter = DocxExporter(classic_template)

    doc_model = DocumentModel(
        blocks=[
            Block(block_type=BlockType.CHAPTER_TITLE, text="Chapter 1: The First Step"),
            Block(block_type=BlockType.HEADING_1, text="1.1 Overview"),
            Block(block_type=BlockType.BODY, text="It was a bright cold day in April, and the clocks were striking thirteen."),
            Block(
                block_type=BlockType.TABLE,
                metadata={"table_data": [["Col 1", "Col 2"], ["Val A", "Val B"]]},
            ),
        ]
    )

    result_path = exporter.export(doc_model, out_file)
    assert os.path.exists(result_path)
    assert os.path.getsize(result_path) > 1000

    # Verify readable in python-docx
    read_back = docx.Document(result_path)
    assert len(read_back.paragraphs) >= 3
    assert len(read_back.tables) >= 1


def test_tabulation_and_table_preservation_roundtrip(tmp_path, classic_template):
    from backend.document.reader import DocxReader

    # 1. Create a DOCX with a table and a tabbed line
    src_file = str(tmp_path / "source_tab_doc.docx")
    src_doc = docx.Document()
    src_doc.add_paragraph("Preamble before table")
    tbl = src_doc.add_table(rows=2, cols=3)
    tbl.cell(0, 0).text = "Item"
    tbl.cell(0, 1).text = "Quantity"
    tbl.cell(0, 2).text = "Cost"
    tbl.cell(1, 0).text = "Notebook"
    tbl.cell(1, 1).text = "10"
    tbl.cell(1, 2).text = "$45.00"

    p_tab = src_doc.add_paragraph()
    p_tab.add_run("Col A")
    p_tab.add_run().add_tab()
    p_tab.add_run("Col B")
    src_doc.save(src_file)

    # 2. Ingest via DocxReader
    reader = DocxReader(src_file)
    blocks = list(reader.stream_raw_blocks())

    # Ensure table cells were NOT emitted as rogue body blocks
    body_texts = [b.text for b in blocks if b.block_type == BlockType.BODY]
    assert "Item" not in body_texts
    assert "Notebook" not in body_texts

    # Ensure table block was extracted with complete grid data
    tbl_blocks = [b for b in blocks if b.block_type == BlockType.TABLE]
    assert len(tbl_blocks) == 1
    table_data = tbl_blocks[0].metadata.get("table_data")
    assert table_data == [
        ["Item", "Quantity", "Cost"],
        ["Notebook", "10", "$45.00"],
    ]

    # Ensure tab character was preserved
    tab_p = [b for b in blocks if "Col A" in b.text][0]
    assert "\t" in tab_p.text

    # 3. Export via DocxExporter
    out_file = str(tmp_path / "exported_tab.docx")
    exporter = DocxExporter(classic_template)
    doc_model = DocumentModel(blocks=blocks)
    exporter.export(doc_model, out_file)

    # 4. Verify exported document has table with values
    exp_doc = docx.Document(out_file)
    assert len(exp_doc.tables) >= 1
    exp_cells = [[c.text for c in r.cells] for r in exp_doc.tables[0].rows]
    assert exp_cells == [
        ["Item", "Quantity", "Cost"],
        ["Notebook", "10", "$45.00"],
    ]


def test_academic_document_alignment_and_header_suppression(tmp_path):
    mgr = TemplateManager()
    academic_template = mgr.get_template("academic")

    blocks = [
        Block(original_index=0, block_type=BlockType.TITLE, text="THE IMPACT OF ARTIFICIAL INTELLIGENCE ON MODERN EDUCATION SYSTEMS: A RESEARCH STUDY"),
        Block(original_index=1, block_type=BlockType.ABSTRACT, text="ABSTRACT"),
        Block(original_index=2, block_type=BlockType.ABSTRACT, text="This research paper investigates AI in education."),
        Block(original_index=3, block_type=BlockType.KEYWORDS, text="Keywords: artificial intelligence, education technology"),
        Block(original_index=4, block_type=BlockType.HEADING_1, text="1. INTRODUCTION"),
        Block(original_index=5, block_type=BlockType.HEADING_2, text="1.1 Background"),
        Block(original_index=6, block_type=BlockType.BODY, text="The rapid advancement of AI has prompted widespread interest."),
        Block(original_index=7, block_type=BlockType.BODY, text="   "),  # Stray empty paragraph
        Block(original_index=8, block_type=BlockType.HEADING_1, text="2. LITERATURE REVIEW"),
        Block(original_index=9, block_type=BlockType.BODY, text="Existing scholarship on AI in education spans several domains."),
        Block(original_index=10, block_type=BlockType.HEADING_1, text="3. METHODOLOGY"),
        Block(original_index=11, block_type=BlockType.HEADING_2, text="3.1 Research Design"),
        Block(original_index=12, block_type=BlockType.BODY, text="This study employs a mixed-methods design."),
    ]

    out_file = str(tmp_path / "academic_perfect_alignment.docx")
    exporter = DocxExporter(academic_template)
    doc_model = DocumentModel(blocks=blocks)
    exporter.export(doc_model, out_file)

    assert os.path.exists(out_file)
    exp_doc = docx.Document(out_file)
    sec = exp_doc.sections[0]

    # First page header MUST be decoupled and empty
    assert sec.different_first_page_header_footer is True
    if sec.first_page_header and sec.first_page_header.paragraphs:
        assert sec.first_page_header.paragraphs[0].text.strip() == ""

    # Running header on continuation pages contains book title
    assert sec.header and len(sec.header.paragraphs) > 0
    assert "THE IMPACT OF ARTIFICIAL INTELLIGENCE" in sec.header.paragraphs[0].text

    # Verify heading properties: keep_with_next is True, styles are Heading 1 and Heading 2
    paras = exp_doc.paragraphs
    h1_paras = [p for p in paras if p.text in ("1. INTRODUCTION", "2. LITERATURE REVIEW", "3. METHODOLOGY")]
    assert len(h1_paras) == 3
    for hp in h1_paras:
        assert hp.paragraph_format.keep_with_next is True
        assert hp.style.name == "Heading 1"

    h2_paras = [p for p in paras if p.text in ("1.1 Background", "3.1 Research Design")]
    assert len(h2_paras) == 2
    for hp in h2_paras:
        assert hp.paragraph_format.keep_with_next is True
        assert hp.style.name == "Heading 2"

    # Verify Author and Affiliation details were inserted on Page 1
    texts = [p.text.strip() for p in paras]
    assert any("Faculty" in t or "Department" in t or "Arun" in t for t in texts[:4])

    # Verify that a page break separates front matter from body content (1. INTRODUCTION)
    intro_idx = next(i for i, p in enumerate(paras) if p.text == "1. INTRODUCTION")
    assert intro_idx >= 5
    # The paragraph immediately preceding 1. INTRODUCTION must contain a page break
    pre_intro = paras[intro_idx - 1]
    has_page_break = bool(pre_intro._p.xpath(".//w:br[@w:type='page']"))
    assert has_page_break is True

    # Verify that stray empty paragraph was suppressed (only 1 empty paragraph, which is the page break)
    empty_paras = [p for p in paras if not p.text.strip()]
    assert len(empty_paras) == 1


def test_page_1_front_matter_isolation_and_page_2_body_start(tmp_path):
    """
    Validates that:
    1. Page 1 contains strictly Front Matter (Title, Author Details & Affiliation, Abstract, Keywords).
    2. Page 1 does NOT contain any body content or section headings.
    3. Content begins strictly on Page 2 via an explicit page break.
    4. End matter (References, advisory callout notes) is formatted with publication-quality styles.
    """
    mgr = TemplateManager()
    template = mgr.get_template("academic")

    raw_blocks = [
        Block(original_index=0, block_type=BlockType.BODY, text="DEEP LEARNING ARCHITECTURES FOR REAL-TIME AUTONOMOUS DECISION SYSTEMS"),
        Block(original_index=1, block_type=BlockType.BODY, text=""),
        Block(original_index=2, block_type=BlockType.BODY, text="ABSTRACT"),
        Block(original_index=3, block_type=BlockType.BODY, text=""),
        Block(original_index=4, block_type=BlockType.BODY, text="This manuscript explores deep reinforcement learning under low-latency constraints."),
        Block(original_index=5, block_type=BlockType.BODY, text=""),
        Block(original_index=6, block_type=BlockType.BODY, text="Keywords: deep learning, reinforcement learning, edge devices, autonomy"),
        Block(original_index=7, block_type=BlockType.BODY, text="1. INTRODUCTION"),
        Block(original_index=8, block_type=BlockType.BODY, text="Autonomous systems require continuous real-time decision loops."),
        Block(original_index=9, block_type=BlockType.BODY, text="2. SYSTEM DESIGN"),
        Block(original_index=10, block_type=BlockType.BODY, text="Our architecture deploys quantized models on edge accelerators."),
        Block(original_index=11, block_type=BlockType.BODY, text="3. CONCLUSION"),
        Block(original_index=12, block_type=BlockType.BODY, text="In conclusion, our method achieves state-of-the-art inference speed."),
        Block(original_index=13, block_type=BlockType.BODY, text="REFERENCES"),
        Block(original_index=14, block_type=BlockType.BODY, text="[1] Y. LeCun, Y. Bengio, and G. Hinton, 'Deep learning,' Nature, vol. 521, 2015."),
        Block(original_index=15, block_type=BlockType.BODY, text="[Note: Advisory note for implementation guidelines.]"),
    ]

    from backend.rules.heading_rules import HeadingHierarchyRule
    classified_blocks = HeadingHierarchyRule.apply(raw_blocks)
    doc_model = DocumentModel(blocks=classified_blocks)

    out_file = str(tmp_path / "page_isolation_test.docx")
    exporter = DocxExporter(template)
    exporter.export(doc_model, out_file)

    assert os.path.exists(out_file)
    exp_doc = docx.Document(out_file)
    paras = exp_doc.paragraphs

    # Verify Page 1 Front Matter items
    texts = [p.text.strip() for p in paras]
    assert "DEEP LEARNING ARCHITECTURES" in texts[0]
    # Author details synthesized
    assert "Dr. Arun Kumar, Ph.D." in texts[1]
    assert "Department of Computer Science" in texts[2]
    assert "Correspondence: arun.kumar@iat.ac.in" in texts[3]
    assert texts[4] == "ABSTRACT"
    assert "This manuscript explores deep reinforcement learning" in texts[5]
    assert "Keywords: deep learning" in texts[6]

    # Verify page break precedes 1. INTRODUCTION
    intro_idx = texts.index("1. INTRODUCTION")
    assert intro_idx == 8
    break_p = paras[intro_idx - 1]
    assert bool(break_p._p.xpath(".//w:br[@w:type='page']")) is True

    # Verify NO body content appears before the page break (items 0-6 are front matter only)
    for t in texts[:7]:
        assert not t.startswith("1. ")
        assert not t.startswith("Autonomous systems")

    # Verify references and callout styling at end of document
    ref_idx = texts.index("REFERENCES")
    assert paras[ref_idx].style.name == "Heading 1"
    assert paras[ref_idx].paragraph_format.keep_with_next is True

    # Citations hanging indent
    cite_p = paras[ref_idx + 1]
    assert cite_p.paragraph_format.left_indent.inches == 0.25
    assert cite_p.paragraph_format.first_line_indent.inches == -0.25

    # Advisory note has shading and border
    note_p = paras[ref_idx + 2]
    assert "[Note:" in note_p.text
    assert note_p.runs[0].font.italic is True
    pPr_xml = note_p._p.pPr.xml
    assert "w:pBdr" in pPr_xml
    assert "w:shd" in pPr_xml



